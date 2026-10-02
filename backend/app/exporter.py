from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def _text(value):
    """Convert strings/dicts/other values safely to text."""
    if value is None:
        return ""

    if isinstance(value, str):
        return value

    if isinstance(value, dict):
        parts = []

        for key in ["title", "owner", "deadline", "decision", "context"]:
            value_part = value.get(key)

            if value_part:
                parts.append(str(value_part))

        return " — ".join(parts)

    return str(value)


def docx_bytes(title, analysis):
    doc = Document()

    # Title
    doc.add_heading(title, level=1)

    # Summary
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        _text(analysis.get("summary", ""))
    )

    # Decisions
    doc.add_heading("Decisions", level=2)

    decisions = analysis.get("decisions", [])

    if decisions:
        for decision in decisions:
            doc.add_paragraph(
                _text(decision),
                style="List Bullet"
            )
    else:
        doc.add_paragraph("No decisions recorded.")

    # Action Items
    doc.add_heading("Action Items", level=2)

    action_items = analysis.get("action_items", [])

    if action_items:
        for item in action_items:
            doc.add_paragraph(
                _text(item),
                style="List Bullet"
            )
    else:
        doc.add_paragraph("No action items recorded.")

    # Risks
    doc.add_heading("Risks", level=2)

    risks = analysis.get("risks", [])

    if risks:
        for risk in risks:
            doc.add_paragraph(
                _text(risk),
                style="List Bullet"
            )
    else:
        doc.add_paragraph("No risks recorded.")

    # Follow-up Questions
    doc.add_heading("Follow-up Questions", level=2)

    questions = analysis.get("follow_up_questions", [])

    if questions:
        for question in questions:
            doc.add_paragraph(
                _text(question),
                style="List Bullet"
            )
    else:
        doc.add_paragraph("No follow-up questions recorded.")

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output


def pdf_bytes(title, analysis):
    output = BytesIO()

    pdf = canvas.Canvas(output, pagesize=letter)

    width, height = letter

    x = 50
    y = height - 50

    def write_line(text, size=10, gap=15):
        nonlocal y

        if y < 50:
            pdf.showPage()
            y = height - 50

        pdf.setFont("Helvetica", size)

        # Basic wrapping
        text = str(text)

        max_chars = 90

        if len(text) <= max_chars:
            pdf.drawString(x, y, text)
            y -= gap
            return

        words = text.split()
        line = ""

        for word in words:
            test_line = line + " " + word if line else word

            if len(test_line) > max_chars:
                pdf.drawString(x, y, line)
                y -= gap

                if y < 50:
                    pdf.showPage()
                    y = height - 50
                    pdf.setFont("Helvetica", size)

                line = word
            else:
                line = test_line

        if line:
            pdf.drawString(x, y, line)
            y -= gap

    # Title
    write_line(title, size=16, gap=22)

    # Summary
    write_line("Summary", size=13, gap=18)
    write_line(
        _text(analysis.get("summary", "")),
        size=10,
        gap=15
    )

    # Decisions
    write_line("Decisions", size=13, gap=18)

    decisions = analysis.get("decisions", [])

    if decisions:
        for decision in decisions:
            write_line(
                "• " + _text(decision),
                size=10,
                gap=15
            )
    else:
        write_line("No decisions recorded.", size=10)

    # Action Items
    write_line("Action Items", size=13, gap=18)

    action_items = analysis.get("action_items", [])

    if action_items:
        for item in action_items:
            write_line(
                "• " + _text(item),
                size=10,
                gap=15
            )
    else:
        write_line("No action items recorded.", size=10)

    # Risks
    write_line("Risks", size=13, gap=18)

    risks = analysis.get("risks", [])

    if risks:
        for risk in risks:
            write_line(
                "• " + _text(risk),
                size=10,
                gap=15
            )
    else:
        write_line("No risks recorded.", size=10)

    # Follow-up Questions
    write_line("Follow-up Questions", size=13, gap=18)

    questions = analysis.get("follow_up_questions", [])

    if questions:
        for question in questions:
            write_line(
                "• " + _text(question),
                size=10,
                gap=15
            )
    else:
        write_line("No follow-up questions recorded.", size=10)

    pdf.save()

    output.seek(0)

    return output